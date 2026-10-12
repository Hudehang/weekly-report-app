# -*- coding: utf-8 -*-
"""
多人周报填报系统（免登录版）
打开即填，自填姓名；本周重复提交自动覆盖（即修改）；
管理员密码查看全部 / 编辑 / 删除单条 / 导出 Excel + Word
运行：streamlit run weekly_report.py
"""
import streamlit as st
import sqlite3
import pandas as pd
from datetime import datetime, date, timedelta
from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ============ 配置区（上线前改这里） ============
ADMIN_PASSWORD = "admin@2026"   # 管理员密码
DB_FILE = "report.db"
# ================================================

def week_key(today: date = None):
    today = today or date.today()
    iso = today.isocalendar()
    monday = today - timedelta(days=today.weekday())
    return f"{iso[0]}-W{iso[1]:02d}", monday.isoformat()

def init_db():
    conn = sqlite3.connect(DB_FILE); c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS week_report(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT, week_key TEXT, week_start TEXT,
        work_this TEXT, work_next TEXT, risk TEXT, updated_at TEXT
    )''')
    conn.commit(); conn.close()

def upsert_report(name, wk, ws, wt, wn, risk):
    conn = sqlite3.connect(DB_FILE); c = conn.cursor()
    c.execute("SELECT id FROM week_report WHERE name=? AND week_key=?", (name, wk))
    row = c.fetchone()
    now = datetime.now().isoformat(timespec="seconds")
    if row:
        c.execute("""UPDATE week_report SET work_this=?,work_next=?,risk=?,updated_at=? WHERE id=?""",
                  (wt, wn, risk, now, row[0]))
    else:
        c.execute("""INSERT INTO week_report(name,week_key,week_start,work_this,work_next,risk,updated_at)
                     VALUES(?,?,?,?,?,?,?)""", (name, wk, ws, wt, wn, risk, now))
    conn.commit(); conn.close()

def get_my_report(name, wk):
    conn = sqlite3.connect(DB_FILE); c = conn.cursor()
    c.execute("SELECT work_this,work_next,risk FROM week_report WHERE name=? AND week_key=?", (name, wk))
    r = c.fetchone(); conn.close(); return r

def get_rows_by_week(wk):
    """带 id 的完整记录，供管理员编辑/删除"""
    conn = sqlite3.connect(DB_FILE)
    df = pd.read_sql("""SELECT id,name,work_this,work_next,risk,updated_at
                        FROM week_report WHERE week_key=? ORDER BY name""",
                     conn, params=(wk,))
    conn.close(); return df

def get_row_by_id(rid):
    conn = sqlite3.connect(DB_FILE); c = conn.cursor()
    c.execute("SELECT id,name,week_key,work_this,work_next,risk,updated_at FROM week_report WHERE id=?", (rid,))
    r = c.fetchone(); conn.close(); return r

def update_report_by_id(rid, name, wt, wn, risk):
    conn = sqlite3.connect(DB_FILE); c = conn.cursor()
    now = datetime.now().isoformat(timespec="seconds")
    c.execute("""UPDATE week_report SET name=?,work_this=?,work_next=?,risk=?,updated_at=?
                 WHERE id=?""", (name, wt, wn, risk, now, rid))
    conn.commit(); conn.close()

def delete_report_by_id(rid):
    conn = sqlite3.connect(DB_FILE); c = conn.cursor()
    c.execute("DELETE FROM week_report WHERE id=?", (rid,))
    conn.commit(); conn.close()

def build_docx(df, week_label):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    h = doc.add_heading(f"周报汇总（{week_label}）", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _, row in df.iterrows():
        doc.add_heading(str(row["name"]), level=2)
        for label, col in [("本周完成工作", "work_this"),
                           ("下周工作计划", "work_next"),
                           ("问题与风险", "risk")]:
            p = doc.add_paragraph()
            p.add_run(f"{label}：").bold = True
            p.add_run(str(row[col]) if row[col] is not None else "")
        doc.add_paragraph()
    buf = BytesIO(); doc.save(buf); buf.seek(0); return buf

init_db()
wk, ws = week_key()

# session_state: 当前正在编辑的记录 id
if "editing_id" not in st.session_state:
    st.session_state["editing_id"] = None

st.title("📝 周报填报系统（免登录）")
st.caption(f"本周：**{wk}**　周一日期：{ws}　打开即填，本周重复提交会覆盖已有内容。")

name = st.text_input("你的姓名", placeholder="输入姓名后填写下方周报")

tab_fill, tab_admin = st.tabs(["📝 填报 / 修改周报", "📊 管理员面板"])

with tab_fill:
    if not name.strip():
        st.info("⬆ 先在上方输入你的姓名")
    else:
        existing = get_my_report(name.strip(), wk)
        default_this = existing[0] if existing else ""
        default_next = existing[1] if existing else ""
        default_risk = existing[2] if existing else ""
        if existing:
            st.info(f"已检测到你本周（{wk}）的周报，可直接修改后重新提交。")
        wt = st.text_area("本周完成工作", value=default_this, height=150, placeholder="可多条，每行一条")
        wn = st.text_area("下周工作计划", value=default_next, height=150)
        risk = st.text_area("存在问题 / 需要协调的风险", value=default_risk, height=100)
        if st.button("提交周报", type="primary"):
            if not (wt.strip() and wn.strip()):
                st.warning("本周工作和下周计划不能为空")
            else:
                upsert_report(name.strip(), wk, ws, wt.strip(), wn.strip(), risk.strip())
                st.success("✅ 周报已提交！本周内可随时回来修改，再次提交会覆盖。")

with tab_admin:
    pwd = st.text_input("管理员密码", type="password")
    if pwd != ADMIN_PASSWORD:
        if pwd: st.error("密码错误")
        st.stop()

    all_weeks = sqlite3.connect(DB_FILE).execute(
        "SELECT DISTINCT week_key FROM week_report ORDER BY week_key DESC").fetchall()
    week_options = [w[0] for w in all_weeks] or [wk]
    sel_week = st.selectbox("选择周次", week_options, index=0)

    df = get_rows_by_week(sel_week)
    submitted_names = sorted(df["name"].tolist()) if not df.empty else []

    c1, c2 = st.columns(2)
    c1.metric("本周已提交人数", len(submitted_names))
    c2.metric("本周记录条数", len(df))
    if submitted_names:
        st.success("已提交人员：" + "、".join(submitted_names))
    else:
        st.info("本周还没有人提交")

    # ---------- 编辑表单（点"编辑"后出现） ----------
    eid = st.session_state["editing_id"]
    if eid is not None:
        row = get_row_by_id(eid)
        if row is None:
            st.session_state["editing_id"] = None
        else:
            rid, rname, rwk, rwt, rwn, rrisk, rupd = row
            st.warning(f"正在编辑：{rname}（{rwk}）　原更新时间：{rupd}")
            e_name = st.text_input("姓名", value=rname, key=f"ename_{rid}")
            e_wt = st.text_area("本周完成工作", value=rwt or "", height=120, key=f"ewt_{rid}")
            e_wn = st.text_area("下周工作计划", value=rwn or "", height=120, key=f"ewn_{rid}")
            e_risk = st.text_area("问题与风险", value=rrisk or "", height=80, key=f"erisk_{rid}")
            b1, b2 = st.columns(2)
            if b1.button("💾 保存修改", key=f"save_{rid}", type="primary"):
                update_report_by_id(rid, e_name.strip(), e_wt.strip(), e_wn.strip(), e_risk.strip())
                st.success("✅ 已保存修改")
                st.session_state["editing_id"] = None
                st.rerun()
            if b2.button("取消", key=f"cancel_{rid}"):
                st.session_state["editing_id"] = None
                st.rerun()
        st.divider()

    # ---------- 逐条记录 + 编辑/删除按钮 ----------
    st.subheader(f"{sel_week} 全部周报（{len(df)} 条）")
    if df.empty:
        st.info("本周还没有人提交")
    else:
        for _, r in df.iterrows():
            rid = r["id"]
            with st.container(border=True):
                cols = st.columns([3, 1, 1])
                cols[0].write(f"**{r['name']}**　·　更新于 {r['updated_at']}")
                cols[0].caption(f"本周：{(r['work_this'] or '')[:60]}　|　下周：{(r['work_next'] or '')[:40]}")
                if cols[1].button("✏️ 编辑", key=f"edit_{rid}"):
                    st.session_state["editing_id"] = rid
                    st.rerun()
                # 删除：点一次变确认态，再点一次才真删
                if st.session_state.get(f"del_confirm_{rid}", False):
                    if cols[2].button("⚠️ 确认删除", key=f"del_yes_{rid}", type="primary"):
                        delete_report_by_id(rid)
                        st.warning(f"已删除 {r['name']} 的周报")
                        st.session_state[f"del_confirm_{rid}"] = False
                        st.rerun()
                else:
                    if cols[2].button("🗑 删除", key=f"del_{rid}"):
                        st.session_state[f"del_confirm_{rid}"] = True
                        st.rerun()

        st.divider()
        # 导出 Excel / Word（导出不含 id 列）
        export_df = df[["name", "work_this", "work_next", "risk", "updated_at"]]
        buf_x = BytesIO()
        with pd.ExcelWriter(buf_x, engine="openpyxl") as xw:
            export_df.to_excel(xw, index=False, sheet_name=sel_week)
        buf_x.seek(0)
        buf_w = build_docx(export_df, sel_week)
        c1, c2 = st.columns(2)
        c1.download_button("📥 导出 Excel", buf_x,
                          file_name=f"周报_{sel_week}.xlsx",
                          mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        c2.download_button("📥 导出 Word", buf_w,
                           file_name=f"周报_{sel_week}.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
